"""The DOCX is the fitted PDF in a second container, set in the same layout.

The regression: every DOCX - five bases and every tailored resume - was three
pages in Microsoft Word while its PDF was one. The writer inherited python-docx's
blank template (1in/1.25in margins, Calibri 11pt, Heading 2 and List Bullet with
their own spacing), so the content the PDF had been fitted to did not fit.

Word is not available to the suite, so these tests pin what made the page count
match when it was measured in Word (2026-10-26, 30 resumes, all one page): the
Spark geometry, faces, sizes and exact line heights, no inherited template
style, and the same words as the PDF.
"""

from __future__ import annotations

import re

import pytest
import yaml
from docx import Document
from docx.enum.text import WD_LINE_SPACING
from docx.shared import Pt
from pypdf import PdfReader

from app.services import resume_library
from app.services import resume_template as tmpl
from app.services.resume_base import BASES_PATH

TWIP = Pt(0.05)
BASES = [b["name"] for b in yaml.safe_load(BASES_PATH.read_text(encoding="utf-8"))["bases"]]


@pytest.fixture(autouse=True)
def isolated_library(tmp_path, monkeypatch):
    """Never write over the real library files while testing."""
    monkeypatch.setattr(resume_library, "BASE_DIR", tmp_path / "bases")


@pytest.fixture(scope="module")
def rendered(tmp_path_factory):
    """Every base, fitted and written by the one path that writes resumes."""
    out = tmp_path_factory.mktemp("bases")
    patch = pytest.MonkeyPatch()
    patch.setattr(resume_library, "BASE_DIR", out)
    try:
        yield {name: resume_library.generate(name) for name in BASES}
    finally:
        patch.undo()


def _words(text: str) -> list[str]:
    text = text.replace(" ", " ").replace("•", " ")
    return text.upper().split()


@pytest.mark.parametrize("base", BASES)
def test_docx_page_is_the_pdf_page(rendered, base):
    for section in Document(str(rendered[base].docx_path)).sections:
        for actual, expected in (
            (section.page_width, tmpl.PAGE_SIZE[0]),
            (section.page_height, tmpl.PAGE_SIZE[1]),
            (section.left_margin, tmpl.LEFT_MARGIN),
            (section.right_margin, tmpl.RIGHT_MARGIN),
            (section.top_margin, tmpl.TOP_MARGIN),
            (section.bottom_margin, tmpl.BOTTOM_MARGIN),
        ):
            # DOCX stores page geometry in whole twips (1/20 pt).
            assert abs(actual - Pt(expected)) <= TWIP, (actual.pt, expected)
        # An empty header still pushes the body down if its distance is
        # larger than the top margin.
        assert section.header_distance <= section.top_margin
        assert section.footer_distance <= section.bottom_margin


@pytest.mark.parametrize("base", BASES)
def test_docx_body_is_the_spark_body(rendered, base):
    normal = Document(str(rendered[base].docx_path)).styles["Normal"]
    assert normal.font.name == tmpl.DOCX_FONT
    assert normal.font.size == Pt(tmpl.BODY_SIZE)
    fmt = normal.paragraph_format
    assert fmt.line_spacing_rule == WD_LINE_SPACING.EXACTLY
    assert fmt.line_spacing == Pt(tmpl.BODY_LEADING)
    assert fmt.space_before == 0 and fmt.space_after == 0


@pytest.mark.parametrize("base", BASES)
def test_docx_inherits_no_template_spacing(rendered, base):
    """No Heading or List style, no font but the body face, no size that is not
    one of the template's, and every line an exact height."""
    allowed_sizes = {None, Pt(tmpl.BODY_SIZE), Pt(tmpl.HEADING_SIZE), Pt(tmpl.NAME_SIZE)}
    doc = Document(str(rendered[base].docx_path))
    for paragraph in doc.paragraphs:
        assert paragraph.style.name == "Normal", paragraph.text
        fmt = paragraph.paragraph_format
        assert fmt.line_spacing_rule in (None, WD_LINE_SPACING.EXACTLY), paragraph.text
        assert (fmt.space_before or 0) <= Pt(0.1), paragraph.text
        assert (fmt.space_after or 0) <= Pt(4), paragraph.text
        for run in paragraph.runs:
            assert run.font.name is None, paragraph.text
            assert run.font.size in allowed_sizes, paragraph.text


@pytest.mark.parametrize("base", BASES)
def test_docx_says_exactly_what_the_fitted_pdf_says(rendered, base):
    """Same words in the same order: the DOCX is not fitted separately, so it
    must not carry anything the one-page PDF dropped, or omit anything it kept."""
    result = rendered[base]
    pdf_text = "\n".join(page.extract_text() for page in PdfReader(result.pdf_path).pages)
    docx_text = "\n".join(p.text for p in Document(str(result.docx_path)).paragraphs)
    assert _words(docx_text) == _words(pdf_text)


def test_docx_bolds_the_same_figures_as_the_pdf(rendered):
    doc = Document(str(rendered["DS_ML_AI"].docx_path))
    bullets = [p for p in doc.paragraphs if p.text.startswith("•\t")]
    assert bullets
    for paragraph in bullets:
        text = paragraph.text[2:]
        if text.startswith("Relevant Coursework") or re.match(r"[^:]{1,40}:", text):
            continue  # coursework and skill lines carry their own emphasis
        bold = [run.text for run in paragraph.runs if run.bold]
        assert bold == [text[a:b] for a, b in tmpl.emphasis_spans(text)]

"""Rendering a cover letter to PDF and DOCX.

Same visual system as the resume — the Times family, the same name treatment,
the same contact line, the same margins — so the two documents arrive looking
like they came from one person. It must still read as a LETTER: no section
rules, no bullets, no two-column date rows. A cover letter that looks like a
second resume is a cover letter nobody reads.

Everything applicant-facing obeys the resume's rules: one page, real extractable
text, no internal identifiers, and the same ATS-safe encoding that keeps bullets
and dashes from decoding as control characters (D-35).
"""

from __future__ import annotations

from datetime import date
from pathlib import Path
from typing import Any
from xml.sax.saxutils import escape

from docx import Document
from docx.shared import Pt
from reportlab.lib.enums import TA_JUSTIFY
from reportlab.lib.styles import ParagraphStyle
from reportlab.platypus import Paragraph, Spacer

from app.config import GENERATED_DIR
from app.models.ingestion import Job
from app.services import resume_template as tmpl
from app.services.candidate_profile import CandidateProfile, load_candidate_profile
from app.services.cover_letter import CoverLetter

#: Where generated letters land. Git-ignored with the resumes — these carry
#: personal data and belong on the machine, not in the repository.
COVER_LETTER_DIR = GENERATED_DIR / "cover_letters"


def _styles() -> dict[str, ParagraphStyle]:
    """Letter styles, built on the resume's typography rather than beside it."""
    base = tmpl.build_styles()
    return {
        "name": base.name,
        "contact": base.contact,
        "meta": ParagraphStyle(
            "LetterMeta",
            fontName=tmpl.SERIF,
            fontSize=tmpl.BODY_SIZE,
            leading=tmpl.BODY_LEADING,
            spaceAfter=2,
        ),
        "body": ParagraphStyle(
            "LetterBody",
            fontName=tmpl.SERIF,
            fontSize=tmpl.BODY_SIZE,
            leading=tmpl.BODY_LEADING + 1.4,
            alignment=TA_JUSTIFY,
            spaceAfter=8,
        ),
    }


def render_pdf(
    letter: CoverLetter,
    job: Job,
    path: Path,
    *,
    profile: CandidateProfile | None = None,
    letter_date: date | None = None,
) -> None:
    """One page, text-based, no decoration beyond the resume's own."""
    profile = profile or load_candidate_profile()
    styles = _styles()
    story: list[Any] = []

    story.append(Paragraph(tmpl.small_caps(profile.name, size=tmpl.NAME_SIZE), styles["name"]))
    story.append(Paragraph(tmpl.contact_paragraph(profile.contact_items), styles["contact"]))
    story.append(Spacer(1, 14))

    story.append(Paragraph(escape(_letter_date(letter_date)), styles["meta"]))
    story.append(Spacer(1, 8))

    company = (job.company_name or "").strip()
    if company:
        story.append(Paragraph(escape(company), styles["meta"]))
    if job.title:
        story.append(Paragraph(escape(f"Re: {job.title}"), styles["meta"]))
    story.append(Spacer(1, 12))

    story.append(Paragraph(escape(letter.salutation), styles["body"]))
    for paragraph in letter.paragraphs:
        story.append(Paragraph(escape(paragraph), styles["body"]))

    story.append(Spacer(1, 6))
    story.append(Paragraph("Best regards,", styles["body"]))
    story.append(Paragraph(escape(profile.name), styles["body"]))

    path.parent.mkdir(parents=True, exist_ok=True)
    document = tmpl.make_document(str(path), profile.name)
    document.build(story)


def render_docx(
    letter: CoverLetter,
    job: Job,
    path: Path,
    *,
    profile: CandidateProfile | None = None,
    letter_date: date | None = None,
) -> None:
    """Ordinary Word paragraphs. No text boxes, no floating elements.

    Those look fine on screen and come apart in a parser, which is the same
    reasoning that keeps the resume to a single flowing column.
    """
    profile = profile or load_candidate_profile()
    document = Document()

    heading = document.add_paragraph()
    run = heading.add_run(profile.name)
    run.bold = True
    run.font.size = Pt(16)

    contact = document.add_paragraph()
    for index, (text, href) in enumerate(profile.contact_items):
        if not text:
            continue
        if index:
            contact.add_run(" | ")
        if href:
            tmpl.add_docx_hyperlink(contact, text, href)
        else:
            contact.add_run(text)

    document.add_paragraph("")
    document.add_paragraph(_letter_date(letter_date))
    if job.company_name:
        document.add_paragraph(job.company_name)
    if job.title:
        document.add_paragraph(f"Re: {job.title}")
    document.add_paragraph("")

    document.add_paragraph(letter.salutation)
    for paragraph in letter.paragraphs:
        document.add_paragraph(paragraph)

    document.add_paragraph("")
    document.add_paragraph("Best regards,")
    document.add_paragraph(profile.name)

    path.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(path))


def output_dir() -> Path:
    """Where letters are written.

    Read through a function rather than the module constant so a test can point
    it at a temporary directory. Without that, running the suite scattered real
    PDFs through the user's own `data/generated/` — harmless but wrong: that
    directory holds documents the user actually sends.
    """
    return COVER_LETTER_DIR


def generate(
    letter: CoverLetter, job: Job, *, cover_letter_version_id: str
) -> tuple[str, str]:
    """Render both formats to disk. Returns (pdf_path, docx_path)."""
    target = output_dir()
    pdf_path = target / f"{cover_letter_version_id}.pdf"
    docx_path = target / f"{cover_letter_version_id}.docx"
    profile = load_candidate_profile()
    render_pdf(letter, job, pdf_path, profile=profile)
    render_docx(letter, job, docx_path, profile=profile)
    return str(pdf_path), str(docx_path)


def filename_for(job: Job, extension: str) -> str:
    """``Jordan_Lee_Cover_Letter_Company_Role.pdf`` for the download header."""
    from app.services.resume_documents import _slug

    name = load_candidate_profile().name.replace(" ", "_")
    return f"{name}_Cover_Letter_{_slug(job.company_name)}_{_slug(job.title)}.{extension}"


def _letter_date(value: date | None) -> str:
    """`August 24, 2026`. Built by hand because `%-d` is POSIX-only and
    `%#d` is Windows-only, and a letter dated "August 04" looks wrong."""
    day = value or date.today()
    return f"{day.strftime('%B')} {day.day}, {day.year}"

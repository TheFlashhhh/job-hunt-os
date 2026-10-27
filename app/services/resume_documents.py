"""Renders a tailored resume: COMPLETE BASE + ACCEPTED PATCHES.

The pipeline is:

    complete role-family base resume
      → job-specific tailoring patches
      → accepted user decisions
      → SafeClaimValidator (already applied when the patch was proposed)
      → complete tailored resume
      → PDF / DOCX

The previous version built the whole document out of accepted recommendations,
so accepting two of them produced a two-bullet page with half of it blank. A
resume must stay a complete resume when the user accepts nothing at all;
tailoring changes emphasis, it does not constitute the document.

Hard constraints, because this is the one place resume *content* is written:

* Every bullet is verbatim text from the Evidence Bank — a safe claim, or a
  recommendation body that already passed ``SafeClaimValidator``. Nothing here
  rewrites, embellishes or invents wording.
* Every bullet has exactly ONE owning entry, resolved from its safe claim.
* No internal identifier reaches the page. Job IDs, Evidence IDs, resume base
  names and Job Hunt OS vocabulary belong in the database and the UI, never in
  a document an employer reads.
"""

from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path
from typing import Any
from xml.sax.saxutils import escape

from docx import Document
from docx.shared import Pt
from reportlab.lib.units import inch
from reportlab.pdfbase.pdfmetrics import stringWidth
from reportlab.platypus import KeepTogether, Paragraph

from app.config import GENERATED_DIR
from app.evidence.service import EvidenceBankService, get_evidence_bank
from app.models.ingestion import Job
from app.models.misc import ResumeRecommendation
from app.services import resume_template as tmpl
from app.services.candidate_profile import load_candidate_profile
from app.services.resume_base import BaseResume, ResumeBaseBuilder, ResumeBullet
from app.services.resume_fitting import (
    ResumeOverflowError,
    count_pages,
    fit_to_one_page,
)
from app.services.role_targeting import RoleTargetProfile

RESUME_DIR = GENERATED_DIR / "resumes"

_SKILLS_PREFIX = "Lead the skills line with: "


@dataclass
class ResumeSection:
    """A rendered block: heading, optional context line, bullets.

    ``title``/``location``/``dates`` are kept alongside the flattened
    ``subheading`` because the Spark layout sets the organization against its
    dates and the role against its location on two aligned rows. Flattening
    them into one string is what produced "OpenText | Software Engineering
    Intern | Bangalore, India | Sep 2023" run together on one line.
    """

    heading: str
    subheading: str
    bullets: list[str]
    kind: str = "EXPERIENCE"
    title: str | None = None
    location: str | None = None
    dates: str | None = None
    context: str | None = None
    #: Public repository link for a project heading. Rendered as a hyperlink.
    url: str | None = None


@dataclass
class ResumeContent:
    name: str
    location: str
    contact_line: str
    education_lines: list[str]
    skill_lines: list[str]
    sections: list[ResumeSection]
    #: The same contact items as (display, href) pairs. Renderers that can
    #: hyperlink use these; `contact_line` stays for anything that cannot.
    contact_items: list[tuple[str, str | None]] = field(default_factory=list)
    #: Structured education and the base summary, for the Spark-style layout.
    education: list[Any] = field(default_factory=list)
    summary: str = ""
    skill_groups: list[tuple[str, str]] = field(default_factory=list)
    #: skill name -> why it qualified (DIRECT_JD, STRONG_INFERENCE, CORE_ROLE,
    #: CORE_PINNED; D-70). Page fitting trims by it. NEVER rendered.
    skill_tiers: dict[str, str] = field(default_factory=dict)
    #: Diagnostics for the UI and tests. NEVER rendered into the document.
    base_name: str = ""
    applied_patches: int = 0
    patch_notes: list[str] = field(default_factory=list)
    #: Present when the resume was built against a RoleTargetProfile.
    targeting: Any | None = None
    #: What the one-page fitting loop removed, and the page count it achieved.
    #: Diagnostics for the UI and tests; NEVER rendered into the document.
    fit_reductions: list[str] = field(default_factory=list)
    page_count: int = 1
    role_archetype: str = ""


def build_content(
    job: Job,
    accepted: list[ResumeRecommendation],
    *,
    resume_base: str,
    bank: EvidenceBankService | None = None,
    target: RoleTargetProfile | None = None,
) -> ResumeContent:
    """Complete base resume, patched by whatever the user accepted.

    With ``target``, the base is POSITIONED for this posting first — different
    bullets, different order, role-appropriate skill groups, a role-specific
    summary — and the accepted patches then apply on top. Both layers draw from
    the same Evidence Bank; targeting reorders, patching emphasises, and neither
    writes prose.
    """
    bank = bank or get_evidence_bank()
    builder = ResumeBaseBuilder(bank)
    base = builder.build(resume_base, target=target)

    applied, notes, skills_override = _apply_patches(base, accepted, builder, target=target)

    # Patches can push the page over its budget: an EMPHASIZE that adds a line
    # is capped within its own entry, but nothing yet re-checks the document as
    # a whole. Without this one real resume spilled onto a second page.
    # User-accepted lines are protected — the weakest UNPATCHED line goes first.
    _enforce_page_budget(
        base, builder, protected={cid for rec in accepted for cid in (rec.safe_claim_ids or [])}
    )

    return content_from_base(
        base,
        bank=bank,
        applied=applied,
        notes=notes,
        skills_override=skills_override,
    )


def build_base_content(
    resume_base: str,
    *,
    target: RoleTargetProfile,
    bank: EvidenceBankService | None = None,
) -> ResumeContent:
    """A base resume as a document, with no job and no accepted patches.

    The Resume Library's path. It is :func:`build_content` minus the two
    job-specific inputs - the posting and the user's accepted changes - and
    nothing else: the same builder selects the same Evidence-Bank claims
    through the same gates. ``target`` is a posting-free profile from
    ``RoleTargetProfileBuilder.for_base``.
    """
    bank = bank or get_evidence_bank()
    base = ResumeBaseBuilder(bank).build(resume_base, target=target)
    return content_from_base(base, bank=bank)


def content_from_base(
    base: BaseResume,
    *,
    bank: EvidenceBankService | None = None,
    applied: int = 0,
    notes: list[str] | None = None,
    skills_override: str | None = None,
) -> ResumeContent:
    """Map a built resume onto the document model both renderers read."""
    profile = load_candidate_profile(bank or get_evidence_bank())
    sections = [
        ResumeSection(
            heading=entry.heading,
            subheading=entry.subheading or (entry.context or ""),
            bullets=[b.text for b in entry.bullets],
            kind=entry.kind,
            title=entry.title,
            location=entry.location,
            dates=entry.dates,
            context=entry.context,
            url=entry.url,
        )
        for entry in base.entries
        if entry.bullets
    ]

    skill_groups = [(g.label, ", ".join(g.skills)) for g in base.skill_groups]
    skill_lines = [group.render() for group in base.skill_groups]
    if skills_override:
        # A skills patch reorders the leading group; it never replaces the set.
        skill_lines = _promote_skills(base, skills_override)

    return ResumeContent(
        name=base.name,
        location=base.location,
        contact_line=base.contact_line,
        contact_items=profile.contact_items,
        education_lines=base.education_lines,
        education=base.education,
        summary=base.summary,
        skill_lines=skill_lines,
        skill_groups=skill_groups,
        skill_tiers={name: tier for g in base.skill_groups for name, tier in g.tiers.items()},
        sections=sections,
        base_name=base.base_name,
        applied_patches=applied,
        patch_notes=list(notes or []),
        targeting=base.targeting,
        role_archetype=base.role_archetype,
    )


def _enforce_page_budget(
    base: BaseResume, builder: ResumeBaseBuilder, *, protected: set[str]
) -> None:
    """Trim back to one page, dropping what the user did not choose.

    A resume that runs to two pages is a different document from the one the
    layout was designed for, so the budget is re-applied after patching rather
    than trusted from before it.
    """
    budget = int(builder.layout.get("max_total_bullets", 11))
    guard = 0
    while base.total_bullets > budget and guard < 50:
        guard += 1
        candidates = [
            (entry, bullet)
            for entry in base.entries
            if len(entry.bullets) > 1
            for bullet in entry.bullets[-1:]
            if bullet.claim_id not in protected
        ]
        if not candidates:
            # Everything left is either a protected choice or the last line
            # holding an entry together. Stop rather than gut the document.
            return
        entry, bullet = candidates[-1]
        entry.bullets.remove(bullet)


def _apply_patches(
    base: BaseResume,
    accepted: list[ResumeRecommendation],
    builder: ResumeBaseBuilder,
    *,
    target: RoleTargetProfile | None = None,
) -> tuple[int, list[str], str | None]:
    """Apply accepted recommendations to the base as targeted edits.

    Each accepted row is one of:

    ``EMPHASIZE``   promote its bullet to the top of its owning entry, adding
                    it only if the base did not already select that claim.
    ``SAFE_REWORD`` replace the wording of the bullet it targets.
    ``DE_EMPHASIZE`` drop a weaker bullet, but never empty an entry.
    ``DO_NOT_ADD``  guardrail only — never becomes resume text.

    Ownership comes from the recommendation's safe claim, not its evidence
    list. A recommendation may cite evidence from several organizations; using
    that list is what rendered one OpenText bullet under both OpenText and
    CHAPA.

    **The stored body is a cache, not the source.** Where a row carries a safe
    claim, the wording is re-resolved through the Evidence Bank at render time
    and the row's own text is discarded. A ``ResumeRecommendation`` can sit in
    the database for weeks; if the Bank's ownership fields or claim text move in
    the meantime, the row is stale, and a stale row must never outrank the Bank.
    This is not hypothetical — rows proposed before an ownership fix carried a
    first-person sentence for a gated claim, and rendering trusted them.
    """
    applied = 0
    notes: list[str] = []
    skills_override: str | None = None

    for rec in accepted:
        body = (rec.body or "").strip()
        if not body:
            continue

        if body.startswith(_SKILLS_PREFIX):
            skills_override = body[len(_SKILLS_PREFIX) :].strip()
            applied += 1
            notes.append("Skills order adjusted for this role.")
            continue

        kind = (rec.kind or "").upper()
        if kind == "DO_NOT_ADD":
            # A guardrail. It exists to keep unsupported keywords OFF the page.
            continue

        if rec.safe_claim_ids:
            current, _ = _current_claim_text(rec, builder)
            if not current:
                notes.append(
                    "Dropped a change whose claim the Evidence Bank no longer permits "
                    "in any wording."
                )
                continue
            body = current

        owner = _owner_for(rec, builder)
        entry = base.entry_for_owner(owner) if owner else None
        if entry is None:
            # No owning entry on this base — the material belongs to an
            # experience this role family does not lead with. Dropping the
            # patch is correct; inventing a section for it is not.
            notes.append(f"Skipped a change with no matching section: {body[:48]}…")
            continue

        if kind == "DE_EMPHASIZE":
            if len(entry.bullets) > 1:
                entry.bullets.pop()
                applied += 1
                notes.append(f"De-emphasised a line under {entry.heading}.")
            continue

        if kind == "SAFE_REWORD":
            reworded = _matching_bullet(entry, rec)
            if reworded is not None:
                reworded.text = body
                applied += 1
                notes.append(f"Reworded a line under {entry.heading}.")
                continue

        # EMPHASIZE (and a SAFE_REWORD with nothing to replace): promote.
        existing = _matching_bullet(entry, rec)
        if existing is not None:
            entry.bullets.remove(existing)
            entry.bullets.insert(0, existing)
        else:
            entry.bullets.insert(
                0,
                ResumeBullet(
                    text=body,
                    claim_id=(rec.safe_claim_ids or [""])[0],
                    owner=entry.organization,
                    strength="STRONG",
                    evidence_ids=list(rec.evidence_ids or []),
                ),
            )
            # Keep the page to one page: adding a line drops the weakest.
            # A targeted build allows a denser entry, so the cap follows the
            # same rule the selection used rather than the untargeted default.
            cap = int(
                builder.layout.get(
                    "max_bullets_per_entry_targeted" if target else "max_bullets_per_entry",
                    3,
                )
            )
            if len(entry.bullets) > cap:
                entry.bullets.pop()
        applied += 1
        notes.append(f"Emphasised a line under {entry.heading}.")

    return applied, notes, skills_override


def _current_claim_text(rec: ResumeRecommendation, builder: ResumeBaseBuilder) -> tuple[str, bool]:
    """The wording the Evidence Bank permits for this row's claim, right now.

    Prefers an ungated claim, exactly as the recommendation service does when
    it first proposes a line. Returns ``("", False)`` when no cited claim has
    any permitted wording — a gated claim with no ``project_level_alternative``
    has none, and the caller must then drop the patch rather than print
    something.
    """
    fallback: tuple[str, bool] | None = None
    for claim_id in rec.safe_claim_ids or []:
        text, downgraded = builder.bank.usable_claim_text(claim_id)
        if not text:
            continue
        if not downgraded:
            return text, False
        if fallback is None:
            fallback = (text, True)
    return fallback or ("", False)


def _owner_for(rec: ResumeRecommendation, builder: ResumeBaseBuilder) -> str | None:
    """The single entry that owns this recommendation's bullet.

    Resolved through the safe claim, which maps to exactly one organization.
    Falling back to the recommendation's flattened evidence list would
    reintroduce the misattribution bug, so it is only used when the row has no
    claim at all — and then only via its FIRST evidence item.
    """
    for claim_id in rec.safe_claim_ids or []:
        owner = builder._claim_owner(claim_id)
        if owner:
            return owner
    for evidence_id in (rec.evidence_ids or [])[:1]:
        record = builder.bank.get_evidence(evidence_id)
        if record is not None and record.organization_or_project:
            return record.organization_or_project
    return None


def _matching_bullet(entry, rec: ResumeRecommendation) -> ResumeBullet | None:
    """The base bullet this recommendation targets, if it is already present."""
    claim_ids = set(rec.safe_claim_ids or [])
    for bullet in entry.bullets:
        if bullet.claim_id and bullet.claim_id in claim_ids:
            return bullet
    return None


def _promote_skills(base: BaseResume, override: str) -> list[str]:
    """Reorder skills so the role-relevant ones lead. Never drops the rest."""
    wanted = [s.strip() for s in override.split(",") if s.strip()]
    lines: list[str] = []
    for group in base.skill_groups:
        promoted = [s for s in group.skills if s in wanted]
        remainder = [s for s in group.skills if s not in wanted]
        ordered = promoted + remainder
        lines.append(f"{group.label}: {', '.join(ordered)}")
    return lines


def _paths(resume_version_id: str) -> tuple[Path, Path]:
    RESUME_DIR.mkdir(parents=True, exist_ok=True)
    return (
        RESUME_DIR / f"{resume_version_id}.pdf",
        RESUME_DIR / f"{resume_version_id}.docx",
    )


# ---------------------------------------------------------------------------
# Rendering — one page, ATS-safe, no decorative columns
# ---------------------------------------------------------------------------


def render_pdf(content: ResumeContent, path: Path | Any) -> None:
    """Render the Spark-style one-page resume.

    Section order follows the reference: summary, education, skills, projects,
    experience. The same function renders a base resume and a job-tailored one
    — tailoring changes which bullets arrive here, never how they are set.
    """
    styles = tmpl.build_styles()
    story: list[Any] = []

    # -- header ----------------------------------------------------------
    story.append(Paragraph(tmpl.small_caps(content.name, size=tmpl.NAME_SIZE), styles.name))
    if content.contact_line:
        # The reference separates contact items with generous spacing rather
        # than pipes. Icons would need a FontAwesome file shipped with the
        # project, so plain text carries it — and extracts more cleanly.
        #
        # The DISPLAYED text stays short (`linkedin.com/in/…`); the hyperlink
        # underneath is the fully-qualified URL. A reader on paper gets the
        # readable form and a reader on screen gets a working link, and the
        # extracted text is the display form either way.
        story.append(Paragraph(tmpl.contact_paragraph(content.contact_items), styles.contact))

    # -- professional summary -------------------------------------------
    if content.summary:
        story.append(tmpl.SectionRule("Professional Summary", tmpl.CONTENT_WIDTH))
        story.append(Paragraph(content.summary, styles.summary))

    # -- education -------------------------------------------------------
    if content.education:
        story.append(tmpl.SectionRule("Education", tmpl.CONTENT_WIDTH))
        for school in content.education:
            story.append(
                tmpl.two_column_row(
                    school.institution,
                    school.dates,
                    (styles.entry_left, styles.entry_right),
                )
            )
            story.append(
                tmpl.two_column_row(
                    school.degree_line,
                    school.location,
                    (styles.sub_left, styles.sub_right),
                )
            )
            if school.coursework:
                story.append(
                    Paragraph(
                        f"<i>Relevant Coursework: {school.coursework}</i>",
                        styles.bullet,
                        bulletText="\u2022",
                    )
                )
    elif content.education_lines:
        story.append(tmpl.SectionRule("Education", tmpl.CONTENT_WIDTH))
        for line in content.education_lines:
            story.append(Paragraph(line, styles.summary))

    # -- skills ----------------------------------------------------------
    if content.skill_groups:
        story.append(tmpl.SectionRule("Skills", tmpl.CONTENT_WIDTH))
        for label, skills in content.skill_groups:
            story.append(tmpl.skill_paragraph(label, skills, styles))
    elif content.skill_lines:
        story.append(tmpl.SectionRule("Skills", tmpl.CONTENT_WIDTH))
        for line in content.skill_lines:
            label, _, rest = line.partition(":")
            story.append(tmpl.skill_paragraph(label.strip(), rest.strip(), styles))

    # -- projects, then experience ---------------------------------------
    projects = [s for s in content.sections if s.kind != "EXPERIENCE"]
    experiences = [s for s in content.sections if s.kind == "EXPERIENCE"]

    if projects:
        story.append(tmpl.SectionRule("Projects", tmpl.CONTENT_WIDTH))
        for section in projects:
            story.extend(_project_block(section, styles))

    if experiences:
        story.append(tmpl.SectionRule("Experience", tmpl.CONTENT_WIDTH))
        for section in experiences:
            story.extend(_experience_block(section, styles))

    doc = tmpl.make_document(path, content.name)
    doc.build(story)


def _project_block(section: ResumeSection, styles: tmpl.Styles) -> list[Any]:
    """One project: a single title line with its year at the right edge.

    The reference keeps name, descriptor and stack on one line rather than
    stacking them, which is why the old two-line "CHAPA / BU Spark" +
    "Affordable Housing... | Fall 2025" looked unlike a resume entry.
    """
    name = escape(section.heading)
    if section.url:
        # The name itself is the link, so it stays clickable even when the
        # short form below does not fit the line.
        name = f'<link href="{escape(section.url, {chr(34): "&quot;"})}">{name}</link>'
    heading = f"<b>{name}</b>"
    if section.title:
        heading += f" \u2014 <b>{escape(section.title)}</b>"

    # The reference keeps a project on ONE title line. The context (course or
    # client) is the optional part, so it is added only when it genuinely fits.
    # Otherwise the heading wraps and the entry stops reading as a resume
    # title. Measured against the real font metrics rather than a length guess.
    # A public repository link takes the context's place: on a printed page it
    # is the one thing a reader cannot recover from the name alone.
    tail = _display_url(section.url) if section.url else section.context
    if tail:
        plain = section.heading
        if section.title:
            plain += f" \u2014 {section.title}"
        plain += f" | {tail}"
        available = tmpl.CONTENT_WIDTH - 2.0 * inch - 6
        if stringWidth(plain, tmpl.SERIF_BOLD, tmpl.BODY_SIZE) <= available:
            if section.url:
                href = escape(section.url, {chr(34): "&quot;"})
                heading += f' | <link href="{href}">{escape(tail)}</link>'
            else:
                heading += f" | <i>{escape(tail)}</i>"

    rows: list[Any] = [
        tmpl.two_column_row(
            heading,
            escape(section.dates or ""),
            (styles.sub_left, styles.entry_right),
        )
    ]
    rows += [tmpl.bullet_paragraph(text, styles) for text in section.bullets]
    return [KeepTogether(rows)] if rows else []


def _display_url(url: str) -> str:
    """``https://github.com/a/b`` -> ``github.com/a/b``, as the header shows links."""
    return url.split("://", 1)[-1].rstrip("/")


def _experience_block(section: ResumeSection, styles: tmpl.Styles) -> list[Any]:
    """One job: employer against dates, then role against location."""
    rows: list[Any] = [
        tmpl.two_column_row(
            escape(section.heading),
            escape(section.dates or ""),
            (styles.entry_left, styles.entry_right),
        )
    ]
    if section.title or section.location:
        rows.append(
            tmpl.two_column_row(
                escape(section.title or ""),
                escape(section.location or ""),
                (styles.sub_left, styles.sub_right),
            )
        )
    rows += [tmpl.bullet_paragraph(text, styles) for text in section.bullets]
    return [KeepTogether(rows)] if rows else []


def render_docx(content: ResumeContent, path: Path) -> None:
    """The same document in the format some portals insist on.

    It carries the same sections in the same order as the PDF. It used to omit
    the professional summary and the coursework line, so the two formats said
    different things about the same candidate — and a portal that takes only
    DOCX got the weaker one.
    """
    doc = Document()

    title = doc.add_paragraph()
    run = title.add_run(content.name)
    run.bold = True
    run.font.size = Pt(16)
    contact = doc.add_paragraph()
    contact.add_run(f"{content.location} | " if content.location else "")
    for index, (text, href) in enumerate(content.contact_items or []):
        if index:
            contact.add_run(" | ")
        if href:
            tmpl.add_docx_hyperlink(contact, text, href)
        else:
            contact.add_run(text)
    if not content.contact_items:
        contact.add_run(content.contact_line)

    if content.summary:
        doc.add_heading("Professional Summary", level=2)
        doc.add_paragraph(content.summary)

    if content.education:
        doc.add_heading("Education", level=2)
        for school in content.education:
            heading = doc.add_paragraph()
            heading.add_run(school.institution).bold = True
            if school.dates:
                heading.add_run(f" — {school.dates}")
            detail = ", ".join(p for p in (school.degree_line, school.location) if p)
            if detail:
                doc.add_paragraph(detail)
            if school.coursework:
                course_line = doc.add_paragraph(f"Relevant Coursework: {school.coursework}")
                course_line.runs[0].italic = True
    elif content.education_lines:
        doc.add_heading("Education", level=2)
        for line in content.education_lines:
            doc.add_paragraph(line)

    if content.skill_lines:
        doc.add_heading("Skills", level=2)
        for line in content.skill_lines:
            doc.add_paragraph(line)

    experiences = [s for s in content.sections if s.kind == "EXPERIENCE"]
    projects = [s for s in content.sections if s.kind != "EXPERIENCE"]

    # Projects before experience, matching the PDF. The two formats are one
    # document in two containers, not two opinions about how to present it.
    for label, group in (("Projects", projects), ("Experience", experiences)):
        if not group:
            continue
        doc.add_heading(label, level=2)
        for section in group:
            heading = doc.add_paragraph()
            heading.add_run(section.heading).bold = True
            if section.url:
                heading.add_run(" | ")
                tmpl.add_docx_hyperlink(heading, _display_url(section.url), section.url)
            if section.subheading:
                sub = doc.add_paragraph(section.subheading)
                sub.runs[0].italic = True
            for bullet in section.bullets:
                doc.add_paragraph(bullet, style="List Bullet")

    doc.save(str(path))


def render_pdf_bytes(content: ResumeContent) -> bytes:
    """The same render, to memory. Used by the one-page fitting loop."""
    import io

    buffer = io.BytesIO()
    render_pdf(content, buffer)
    return buffer.getvalue()


def generate(
    job: Job,
    accepted: list[ResumeRecommendation],
    *,
    resume_base: str,
    resume_version_id: str,
    bank: EvidenceBankService | None = None,
    target: RoleTargetProfile | None = None,
) -> tuple[str, str]:
    """Render both formats to disk. Returns (pdf_path, docx_path) as strings.

    The content is FITTED before it is written: rendered, counted, and reduced
    until it is one page. The entry and bullet budgets are a starting point, not
    the invariant — they were satisfied by a two-page real resume that made
    this necessary (D-52).
    """
    content = build_content(job, accepted, resume_base=resume_base, bank=bank, target=target)
    pdf_path, docx_path = _paths(resume_version_id)
    write_fitted(content, pdf_path, docx_path)
    return str(pdf_path), str(docx_path)


def write_fitted(content: ResumeContent, pdf_path: Path, docx_path: Path) -> ResumeContent:
    """Fit to one page, render both formats, and prove the file is one page.

    The single place a resume reaches disk - a job resume and a library base
    both come through here, so neither can skip the fitting loop.
    """
    # Fit on the PDF, then render the DOCX from the same reduced content, so
    # the two documents say the same thing.
    fit = fit_to_one_page(content, render_pdf_bytes)
    content.fit_reductions = list(fit.reductions)
    content.page_count = fit.pages

    pdf_path.parent.mkdir(parents=True, exist_ok=True)
    render_pdf(content, pdf_path)
    render_docx(content, docx_path)

    # Belt and braces: the file on disk is what the employer opens, and it is
    # cheap to confirm the thing that was just proved in memory.
    written = count_pages(pdf_path.read_bytes())
    if written > 1:
        raise ResumeOverflowError(
            f"The rendered PDF is {written} pages after fitting. Nothing was served."
        )
    return content


def filename_for(job: Job, extension: str) -> str:
    """``Jordan_Lee_Company_Role.pdf`` — for the Content-Disposition header.

    The filename is the one place a job reference is acceptable: it helps the
    user manage downloads and is not part of the document an employer reads.
    """
    bank = get_evidence_bank()
    name = str(bank.candidate.get("name", "Resume")).replace(" ", "_")
    return f"{name}_{_slug(job.company_name)}_{_slug(job.title)}.{extension}"


def _slug(text: str) -> str:
    keep = [c if c.isalnum() else "_" for c in (text or "")]
    slug = "".join(keep)
    while "__" in slug:
        slug = slug.replace("__", "_")
    return slug.strip("_")[:60] or "Role"
